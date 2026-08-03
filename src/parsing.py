"""课件解析：pptx / docx / pdf / md / txt → 带页码的文本单元。"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import List

from src.config import SUPPORTED_EXTENSIONS


@dataclass
class PageUnit:
    """一个父单元（幻灯片或页），供分块与 parent retrieval 使用。"""

    page: int
    text: str


def parse_file(path: Path) -> List[PageUnit]:
    """按扩展名解析文件，返回非空页/幻灯片列表。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {suffix}，支持: {sorted(SUPPORTED_EXTENSIONS)}")
    if suffix == ".pptx":
        units = _parse_pptx(path)
    elif suffix == ".docx":
        units = _parse_docx(path)
    elif suffix == ".pdf":
        units = _parse_pdf(path)
    elif suffix in {".md", ".txt"}:
        units = _parse_plain(path)
    else:
        raise ValueError(f"不支持的文件类型: {suffix}")
    units = [u for u in units if u.text.strip()]
    if not units:
        raise ValueError(f"未能从文件中提取到有效文本: {path.name}")
    return units


def _parse_pptx(path: Path) -> List[PageUnit]:
    from pptx import Presentation

    prs = Presentation(str(path))
    units: List[PageUnit] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_texts(shape))
        text = "\n".join(p for p in parts if p.strip()).strip()
        units.append(PageUnit(page=idx, text=text))
    return units


def _shape_texts(shape) -> List[str]:
    """递归提取形状中的文本（含组合形状）。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts: List[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
        return texts
    if getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if line:
                texts.append(line)
    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))
    return texts


def _parse_docx(path: Path) -> List[PageUnit]:
    from docx import Document

    doc = Document(str(path))
    # docx 无分页符弱，按段落聚合后按固定段落数切「伪页」
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    if not paragraphs:
        return []
    page_size = 12
    units: List[PageUnit] = []
    for i in range(0, len(paragraphs), page_size):
        chunk = paragraphs[i : i + page_size]
        units.append(PageUnit(page=len(units) + 1, text="\n".join(chunk)))
    return units


def _parse_pdf(path: Path) -> List[PageUnit]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units: List[PageUnit] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        units.append(PageUnit(page=idx, text=text))
    return units


def _parse_plain(path: Path) -> List[PageUnit]:
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    # 按空行分段为「页」
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    if not blocks:
        return [PageUnit(page=1, text=text)]
    return [PageUnit(page=i, text=b) for i, b in enumerate(blocks, start=1)]
